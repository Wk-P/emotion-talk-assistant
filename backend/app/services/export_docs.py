"""Human-readable exports of admin conversation data (Markdown, TXT, Word,
PDF). JSON stays the machine-readable format and is built in app/api/admin.py.

All formats share one layout: an export header, then per conversation its
participant label / language / times, the transcript, and any records the
participant chose to save. Labels follow the admin's UI language; times are
shown in the admin's own time zone (`tz_offset`, as JS getTimezoneOffset()).
"""

import io
import re
from datetime import datetime, timedelta
from typing import Any
from xml.sax.saxutils import escape

from app.schemas.admin import AdminSessionExport

LABELS: dict[str, dict[str, str]] = {
    "zh": {
        "title": "对话数据导出",
        "exported_at": "导出时间",
        "count": "对话数",
        "conversation": "对话",
        "participant": "参与者",
        "language": "对话语言",
        "started": "开始时间",
        "ended": "结束时间",
        "not_ended": "未结束",
        "messages": "对话内容",
        "no_messages": "（没有保存消息内容）",
        "records": "用户保存的记录",
        "user": "用户",
        "assistant": "助手",
        "system": "系统",
        "lang_zh": "中文",
        "lang_ko": "韩语",
    },
    "ko": {
        "title": "대화 데이터 내보내기",
        "exported_at": "내보낸 시간",
        "count": "대화 수",
        "conversation": "대화",
        "participant": "참여자",
        "language": "대화 언어",
        "started": "시작 시간",
        "ended": "종료 시간",
        "not_ended": "종료되지 않음",
        "messages": "대화 내용",
        "no_messages": "(저장된 메시지가 없습니다)",
        "records": "사용자가 저장한 기록",
        "user": "사용자",
        "assistant": "도우미",
        "system": "시스템",
        "lang_zh": "중국어",
        "lang_ko": "한국어",
    },
    # English is an interface language only (conversations are zh/ko).
    "en": {
        "title": "Conversation data export",
        "exported_at": "Exported at",
        "count": "Conversations",
        "conversation": "Conversation",
        "participant": "Participant",
        "language": "Conversation language",
        "started": "Started",
        "ended": "Ended",
        "not_ended": "Not ended",
        "messages": "Messages",
        "no_messages": "(No messages were kept)",
        "records": "Records saved by the user",
        "user": "User",
        "assistant": "Assistant",
        "system": "System",
        "lang_zh": "Chinese",
        "lang_ko": "Korean",
    },
}

# Punctuation between a label and its value, and around times: full-width
# for Chinese, normal spacing for Korean and English.
PUNCT = {
    "zh": {"c": "：", "l": "（", "r": "）", "h": "【{}】"},
    "ko": {"c": ": ", "l": " (", "r": ")", "h": "[{}]"},
    "en": {"c": ": ", "l": " (", "r": ")", "h": "[{}]"},
}

# Same labels as the frontend (src/utils/fieldLabels.ts, src/i18n records.types).
RECORD_TYPES = {
    "zh": {
        "situation_emotion_behavior": "情境 · 感受 · 应对",
        "cause_interpretation": "原因与理解",
        "value_goal": "在意的事与目标",
        "self_kindness": "对自己友善的话",
        "self_encouragement": "鼓励自己的话",
        "recovery_plan": "行动计划",
        "weekly_reflection": "每周回顾",
    },
    "ko": {
        "situation_emotion_behavior": "상황 · 감정 · 대처",
        "cause_interpretation": "원인과 이해",
        "value_goal": "소중한 것과 목표",
        "self_kindness": "나에게 건네는 따뜻한 말",
        "self_encouragement": "나를 격려하는 말",
        "recovery_plan": "실행 계획",
        "weekly_reflection": "주간 돌아보기",
    },
    "en": {
        "situation_emotion_behavior": "Situation · feeling · response",
        "cause_interpretation": "Causes and understanding",
        "value_goal": "What matters and goals",
        "self_kindness": "Kind words to myself",
        "self_encouragement": "Words of encouragement",
        "recovery_plan": "Action plan",
        "weekly_reflection": "Weekly review",
    },
}
FIELDS = {
    "zh": {
        "situation": "发生了什么",
        "emotion": "当时的感受",
        "behavior": "当时怎么应对",
        "action": "先做的行动",
        "when": "什么时候做",
        "support": "需要的帮助",
        "backup": "做不到时的备选方案",
        "text": "内容",
    },
    "ko": {
        "situation": "무슨 일이 있었는지",
        "emotion": "그때의 감정",
        "behavior": "그때 어떻게 대처했는지",
        "action": "먼저 할 행동",
        "when": "언제 할지",
        "support": "필요한 도움",
        "backup": "어려울 때의 대안",
        "text": "내용",
    },
    "en": {
        "situation": "What happened",
        "emotion": "How I felt",
        "behavior": "How I responded",
        "action": "First action",
        "when": "When",
        "support": "Help needed",
        "backup": "Backup plan",
        "text": "Content",
    },
}


def _fmt_time(iso: str | None, tz_offset: int) -> str:
    if not iso:
        return ""
    dt = datetime.fromisoformat(iso)
    if dt.tzinfo is not None:
        dt = dt.replace(tzinfo=None) - (dt.utcoffset() or timedelta())
    # Stored times are UTC; getTimezoneOffset() is "UTC minus local".
    return (dt - timedelta(minutes=tz_offset)).strftime("%Y-%m-%d %H:%M")


def _field_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "、".join(t for t in (_field_text(v) for v in value) if t)
    if isinstance(value, dict):
        return "；".join(t for t in (_field_text(v) for v in value.values()) if t)
    return str(value)


def _document(sessions: list[AdminSessionExport], lang: str, tz_offset: int) -> dict[str, Any]:
    """The export as plain data, shared by every renderer below."""

    L = LABELS[lang]
    now = (datetime.utcnow() - timedelta(minutes=tz_offset)).strftime("%Y-%m-%d %H:%M")
    convs = []
    for i, s in enumerate(sessions, 1):
        who = s.participant_label
        meta = [
            (L["participant"], who),
            (L["language"], L.get(f"lang_{s.language.value}", s.language.value)),
            (L["started"], _fmt_time(s.created_at, tz_offset)),
            (L["ended"], _fmt_time(s.ended_at, tz_offset) or L["not_ended"]),
        ]
        messages = [(L.get(m.role, m.role), _fmt_time(m.created_at, tz_offset), m.content) for m in s.messages]
        records = []
        for r in s.records:
            lines = [
                (FIELDS[lang].get(k, k.replace("_", " ")), _field_text(v))
                for k, v in r.payload.items()
                if _field_text(v)
            ]
            records.append((RECORD_TYPES[lang].get(r.record_type, r.record_type), _fmt_time(r.created_at, tz_offset), lines))
        convs.append({"heading": f"{L['conversation']} {i} · {who}", "meta": meta, "messages": messages, "records": records})
    return {"title": L["title"], "meta": [(L["exported_at"], now), (L["count"], str(len(sessions)))], "convs": convs, "L": L}


def to_markdown(sessions: list[AdminSessionExport], lang: str, tz_offset: int) -> str:
    P = PUNCT[lang]
    d = _document(sessions, lang, tz_offset)
    L = d["L"]
    out = [f"# {d['title']}", ""]
    out += [f"- {k}{P['c']}{v}" for k, v in d["meta"]]
    for c in d["convs"]:
        out += ["", "---", "", f"## {c['heading']}", ""]
        out += [f"- **{k}**{P['c']}{v}" for k, v in c["meta"]]
        out += ["", f"### {L['messages']}", ""]
        if not c["messages"]:
            out.append(L["no_messages"])
        for role, time, content in c["messages"]:
            body = content.replace("\n", "  \n")
            out += [f"**{role}**{P['l']}{time}{P['r']}  ", body, ""]
        if c["records"]:
            out += [f"### {L['records']}", ""]
            for rtype, time, lines in c["records"]:
                out.append(f"- **{rtype}**{P['l']}{time}{P['r']}")
                out += [f"  - {k}{P['c']}{v}" for k, v in lines]
    return "\n".join(out).rstrip() + "\n"


def to_text(sessions: list[AdminSessionExport], lang: str, tz_offset: int) -> str:
    P = PUNCT[lang]
    d = _document(sessions, lang, tz_offset)
    L = d["L"]
    out = [d["title"], "=" * 40]
    out += [f"{k}{P['c']}{v}" for k, v in d["meta"]]
    for c in d["convs"]:
        out += ["", "-" * 40, c["heading"], "-" * 40]
        out += [f"{k}{P['c']}{v}" for k, v in c["meta"]]
        out += ["", P["h"].format(L["messages"])]
        if not c["messages"]:
            out.append(L["no_messages"])
        for role, time, content in c["messages"]:
            out += [f"[{time}] {role}{P['c']}", content, ""]
        if c["records"]:
            out.append(P["h"].format(L["records"]))
            for rtype, time, lines in c["records"]:
                out.append(f"· {rtype}{P['l']}{time}{P['r']}")
                out += [f"    {k}{P['c']}{v}" for k, v in lines]
    return "\n".join(out).rstrip() + "\n"


def to_docx(sessions: list[AdminSessionExport], lang: str, tz_offset: int) -> bytes:
    P = PUNCT[lang]
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = _document(sessions, lang, tz_offset)
    L = d["L"]
    east_asian = "Malgun Gothic" if lang == "ko" else "Microsoft YaHei"

    doc = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = east_asian
        style.element.rPr.rFonts.set(qn("w:eastAsia"), east_asian)
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(d["title"], level=0)
    for k, v in d["meta"]:
        doc.add_paragraph(f"{k}{P['c']}{v}")
    for c in d["convs"]:
        doc.add_heading(c["heading"], level=1)
        for k, v in c["meta"]:
            p = doc.add_paragraph()
            p.add_run(f"{k}{P['c']}").bold = True
            p.add_run(v)
        doc.add_heading(L["messages"], level=2)
        if not c["messages"]:
            doc.add_paragraph(L["no_messages"])
        for role, time, content in c["messages"]:
            p = doc.add_paragraph()
            who = p.add_run(f"{role}")
            who.bold = True
            if role == L["user"]:
                who.font.color.rgb = RGBColor(0x6C, 0x5C, 0xE7)
            t = p.add_run(f"  {time}")
            t.font.size = Pt(8.5)
            t.font.color.rgb = RGBColor(0x71, 0x75, 0x8C)
            p.add_run("\n" + content)
        if c["records"]:
            doc.add_heading(L["records"], level=2)
            for rtype, time, lines in c["records"]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(rtype).bold = True
                p.add_run(f"{P['l']}{time}{P['r']}")
                for k, v in lines:
                    doc.add_paragraph(f"{k}{P['c']}{v}", style="List Bullet 2")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# reportlab's built-in CID fonts — no font files needed. The Chinese one has
# no Hangul, so Korean runs are switched to the Korean font inline.
_ZH_FONT = "STSong-Light"
_KO_FONT = "HYSMyeongJo-Medium"
_LATIN_FONT = "Helvetica"
# Hangul → Korean font; Latin letters, digits, spaces and Latin-1 symbols
# such as "·" → Helvetica (the CID fonts squeeze spaces and lack "·");
# everything else (Chinese, CJK punctuation) stays in the Chinese font.
_RUNS = re.compile("([ᄀ-ᇿ㄰-㆏가-힯]+)|([\x20-\x7e -ÿ]+)")
_fonts_ready = False


def _pdf_text(text: str) -> str:
    lines = []
    for line in text.split("\n"):
        pos, parts = 0, []
        for m in _RUNS.finditer(line):
            parts.append(escape(line[pos : m.start()]))
            font = _KO_FONT if m.group(1) else _LATIN_FONT
            parts.append(f'<font name="{font}">{escape(m.group(0))}</font>')
            pos = m.end()
        parts.append(escape(line[pos:]))
        lines.append("".join(parts))
    return "<br/>".join(lines)


def to_pdf(sessions: list[AdminSessionExport], lang: str, tz_offset: int) -> bytes:
    P = PUNCT[lang]
    global _fonts_ready
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    if not _fonts_ready:
        pdfmetrics.registerFont(UnicodeCIDFont(_ZH_FONT))
        pdfmetrics.registerFont(UnicodeCIDFont(_KO_FONT))
        _fonts_ready = True

    d = _document(sessions, lang, tz_offset)
    L = d["L"]
    accent = colors.HexColor("#6C5CE7")
    muted = colors.HexColor("#71758C")
    base = ParagraphStyle("base", fontName=_ZH_FONT, fontSize=10, leading=15, wordWrap="CJK")
    styles = {
        "title": ParagraphStyle("title", parent=base, fontSize=18, leading=24, spaceAfter=6),
        "h1": ParagraphStyle("h1", parent=base, fontSize=13.5, leading=19, textColor=accent, spaceBefore=10, spaceAfter=4),
        "h2": ParagraphStyle("h2", parent=base, fontSize=11, leading=16, spaceBefore=8, spaceAfter=3),
        "meta": ParagraphStyle("meta", parent=base, fontSize=9.5, textColor=muted),
        "who": ParagraphStyle("who", parent=base, fontSize=9, textColor=muted, spaceBefore=5),
        "msg": ParagraphStyle("msg", parent=base, leftIndent=8),
        "rec": ParagraphStyle("rec", parent=base, leftIndent=8, spaceBefore=3),
    }

    story = [Paragraph(_pdf_text(d["title"]), styles["title"])]
    story += [Paragraph(_pdf_text(f"{k}{P['c']}{v}"), styles["meta"]) for k, v in d["meta"]]
    for c in d["convs"]:
        story += [Spacer(1, 6), HRFlowable(width="100%", color=colors.HexColor("#E6E3F0")), Paragraph(_pdf_text(c["heading"]), styles["h1"])]
        story += [Paragraph(_pdf_text(f"{k}{P['c']}{v}"), styles["meta"]) for k, v in c["meta"]]
        story.append(Paragraph(_pdf_text(L["messages"]), styles["h2"]))
        if not c["messages"]:
            story.append(Paragraph(_pdf_text(L["no_messages"]), styles["msg"]))
        for role, time, content in c["messages"]:
            color = "#6C5CE7" if role == L["user"] else "#1F2333"
            story.append(Paragraph(f'<font color="{color}"><b>{_pdf_text(role)}</b></font>  {_pdf_text(time)}', styles["who"]))
            story.append(Paragraph(_pdf_text(content), styles["msg"]))
        if c["records"]:
            story.append(Paragraph(_pdf_text(L["records"]), styles["h2"]))
            for rtype, time, lines in c["records"]:
                story.append(Paragraph(f"<b>{_pdf_text(rtype)}</b>{P['l']}{_pdf_text(time)}{P['r']}", styles["rec"]))
                story += [Paragraph(_pdf_text(f"{k}{P['c']}{v}"), styles["msg"]) for k, v in lines]

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=d["title"]
    ).build(story)
    return buf.getvalue()
